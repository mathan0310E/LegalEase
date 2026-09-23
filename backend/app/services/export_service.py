import io
from typing import Dict, Any, List
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether, HRFlowable
)
from reportlab.pdfgen import canvas
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from app.models.document import Document


class NumberedCanvas(canvas.Canvas):
    """Two-pass canvas to dynamically compute and draw total page count and professional headers/footers."""
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []
        self.header_text = "CONFIDENTIAL LEGAL DRAFT"
        self.org_name = "LegalEase AI"
        self.footer_text = "Generated via LegalEase AI • For Informational Purposes"

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_header_footer(num_pages)
            super().showPage()
        super().save()

    def draw_header_footer(self, page_count):
        self.saveState()
        self.setFont("Helvetica-Bold", 8)
        self.setFillColor(colors.HexColor("#0f172a"))

        # Header
        self.drawString(54, 750, self.org_name.upper())
        self.setFont("Helvetica-Oblique", 8)
        self.setFillColor(colors.HexColor("#64748b"))
        self.drawRightString(558, 750, self.header_text)
        self.setStrokeColor(colors.HexColor("#cbd5e1"))
        self.setLineWidth(0.5)
        self.line(54, 744, 558, 744)

        # Footer
        self.line(54, 45, 558, 45)
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#64748b"))
        self.drawString(54, 32, self.footer_text)
        page_str = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(558, 32, page_str)
        self.restoreState()


class ExportService:
    @staticmethod
    def generate_pdf(doc: Document) -> bytes:
        """Generates an executive, beautifully typeset legal PDF."""
        buffer = io.BytesIO()
        pdf = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            leftMargin=54,
            rightMargin=54,
            topMargin=64,
            bottomMargin=64
        )

        branding: Dict[str, Any] = doc.branding_config or {}
        structured: Dict[str, Any] = doc.structured_content or {}
        sections: List[Dict[str, Any]] = structured.get("sections", [])
        terms: List[Dict[str, Any]] = structured.get("important_terms", [])
        parties: List[Dict[str, Any]] = structured.get("parties", [])

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'DocTitle',
            parent=styles['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=16,
            leading=20,
            textColor=colors.HexColor("#0f172a"),
            alignment=1,  # Center
            spaceAfter=14
        )
        subtitle_style = ParagraphStyle(
            'DocSubtitle',
            fontName='Helvetica',
            fontSize=10,
            leading=13,
            textColor=colors.HexColor("#475569"),
            alignment=1,
            spaceAfter=16
        )
        h2_style = ParagraphStyle(
            'DocH2',
            parent=styles['Heading2'],
            fontName='Helvetica-Bold',
            fontSize=11,
            leading=15,
            textColor=colors.HexColor("#1e293b"),
            spaceBefore=12,
            spaceAfter=6,
            keepWithNext=True
        )
        body_style = ParagraphStyle(
            'DocBody',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9.5,
            leading=14,
            textColor=colors.HexColor("#334155"),
            spaceAfter=8,
            alignment=4  # Justified
        )
        table_cell_bold = ParagraphStyle(
            'TableCellBold',
            fontName='Helvetica-Bold',
            fontSize=8.5,
            leading=11,
            textColor=colors.HexColor("#0f172a")
        )
        table_cell = ParagraphStyle(
            'TableCell',
            fontName='Helvetica',
            fontSize=8.5,
            leading=11,
            textColor=colors.HexColor("#334155")
        )
        disclaimer_style = ParagraphStyle(
            'DocDisclaimer',
            fontName='Helvetica-Oblique',
            fontSize=7.5,
            leading=10,
            textColor=colors.HexColor("#64748b"),
            alignment=1
        )

        story = []

        # Document Title
        story.append(Paragraph(doc.title.upper(), title_style))
        doc_type_label = doc.document_type.replace('_', ' ').title()
        eff_date = structured.get("effective_date", "Date of Execution")
        story.append(Paragraph(f"<b>Type:</b> {doc_type_label} &nbsp;&nbsp;|&nbsp;&nbsp; <b>Effective Date:</b> {eff_date}", subtitle_style))
        story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#0f172a"), spaceBefore=0, spaceAfter=14))

        # Important Terms Summary Box (if present)
        if terms:
            story.append(Paragraph("KEY TRANSACTION TERMS", h2_style))
            term_table_data = [[Paragraph("Term / Parameter", table_cell_bold), Paragraph("Agreed Specification", table_cell_bold)]]
            for t in terms[:6]:
                term_table_data.append([
                    Paragraph(str(t.get("term", "")), table_cell_bold),
                    Paragraph(str(t.get("value", "")), table_cell)
                ])
            term_table = Table(term_table_data, colWidths=[180, 324])
            term_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#f1f5f9")),
                ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
                ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
                ('TOPPADDING', (0, 0), (-1, -1), 4),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ]))
            story.append(term_table)
            story.append(Spacer(1, 14))

        # Sections
        for sec in sorted(sections, key=lambda s: s.get("order", 0)):
            heading = sec.get("heading", "")
            content = sec.get("content", "")
            story.append(Paragraph(heading, h2_style))
            # Split paragraphs if multiple lines
            for paragraph_text in content.split("\n\n"):
                if paragraph_text.strip():
                    story.append(Paragraph(paragraph_text.strip(), body_style))
            story.append(Spacer(1, 4))

        story.append(Spacer(1, 16))

        # Signature Area
        story.append(Paragraph("IN WITNESS WHEREOF, the Parties have executed this Agreement as of the Effective Date.", body_style))
        story.append(Spacer(1, 12))

        party1_name = parties[0].get("name", "First Party") if len(parties) > 0 else "Party A"
        party2_name = parties[1].get("name", "Second Party") if len(parties) > 1 else "Party B"

        sig_data = [
            [Paragraph(f"<b>For and on behalf of:</b><br/><b>{party1_name}</b>", table_cell_bold),
             Paragraph(f"<b>For and on behalf of:</b><br/><b>{party2_name}</b>", table_cell_bold)],
            [Spacer(1, 30), Spacer(1, 30)],
            [Paragraph("By: _______________________________", table_cell), Paragraph("By: _______________________________", table_cell)],
            [Paragraph("Name: _____________________________", table_cell), Paragraph("Name: _____________________________", table_cell)],
            [Paragraph("Title: ____________________________", table_cell), Paragraph("Title: ____________________________", table_cell)],
            [Paragraph("Date: _____________________________", table_cell), Paragraph("Date: _____________________________", table_cell)],
        ]
        sig_table = Table(sig_data, colWidths=[250, 254])
        sig_table.setStyle(TableStyle([
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('TOPPADDING', (0, 0), (-1, -1), 2),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
        ]))
        story.append(KeepTogether(sig_table))

        story.append(Spacer(1, 20))
        # Legal Disclaimer
        disclaimer_text = structured.get("disclaimer", (
            "LegalEase provides AI-generated document drafts for informational and drafting purposes. "
            "These documents do not constitute legal advice and should be reviewed by a qualified legal professional before use."
        ))
        story.append(Paragraph(disclaimer_text, disclaimer_style))

        # Custom canvas with branding
        def make_canvas(*args, **kwargs):
            c = NumberedCanvas(*args, **kwargs)
            if branding.get("org_name"):
                c.org_name = branding["org_name"]
            if branding.get("header_text"):
                c.header_text = branding["header_text"]
            if branding.get("footer_text"):
                c.footer_text = branding["footer_text"]
            return c

        pdf.build(story, canvasmaker=make_canvas)
        buffer.seek(0)
        return buffer.getvalue()

    @staticmethod
    def generate_docx(doc: Document) -> bytes:
        """Generates a professional Microsoft Word (.docx) document."""
        document = docx.Document()
        branding: Dict[str, Any] = doc.branding_config or {}
        structured: Dict[str, Any] = doc.structured_content or {}
        sections: List[Dict[str, Any]] = structured.get("sections", [])
        terms: List[Dict[str, Any]] = structured.get("important_terms", [])
        parties: List[Dict[str, Any]] = structured.get("parties", [])

        # Configure Header
        header = document.sections[0].header
        header_p = header.paragraphs[0]
        header_p.text = f"{branding.get('org_name', 'LegalEase AI')} | {branding.get('header_text', 'CONFIDENTIAL LEGAL DRAFT')}"
        header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        header_p.style.font.size = Pt(8.5)
        header_p.style.font.color.rgb = RGBColor(100, 116, 139)

        # Title
        title_p = document.add_paragraph()
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title_p.add_run(doc.title.upper())
        title_run.bold = True
        title_run.font.size = Pt(16)
        title_run.font.color.rgb = RGBColor(15, 23, 42)

        # Subtitle
        sub_p = document.add_paragraph()
        sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub_run = sub_p.add_run(f"Document Type: {doc.document_type.replace('_', ' ').title()}  |  Effective Date: {structured.get('effective_date', 'As Executed')}")
        sub_run.font.size = Pt(10)
        sub_run.font.color.rgb = RGBColor(71, 85, 105)

        document.add_paragraph().paragraph_format.space_after = Pt(8)

        # Terms Table
        if terms:
            table = document.add_table(rows=1, cols=2)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = "Contract Term"
            hdr_cells[1].text = "Agreed Parameter"
            hdr_cells[0].paragraphs[0].runs[0].bold = True
            hdr_cells[1].paragraphs[0].runs[0].bold = True

            for item in terms[:8]:
                row_cells = table.add_row().cells
                row_cells[0].text = str(item.get("term", ""))
                row_cells[0].paragraphs[0].runs[0].bold = True
                row_cells[1].text = str(item.get("value", ""))

            document.add_paragraph().paragraph_format.space_after = Pt(12)

        # Document Clauses
        for sec in sorted(sections, key=lambda s: s.get("order", 0)):
            h = document.add_heading(sec.get("heading", ""), level=2)
            h.style.font.size = Pt(12)
            h.style.font.color.rgb = RGBColor(15, 23, 42)

            for p_text in sec.get("content", "").split("\n\n"):
                if p_text.strip():
                    p = document.add_paragraph(p_text.strip())
                    p.paragraph_format.space_after = Pt(6)
                    p.paragraph_format.line_spacing = 1.15
                    p.style.font.size = Pt(10.5)

        # Signatures
        document.add_paragraph().paragraph_format.space_after = Pt(14)
        document.add_paragraph("IN WITNESS WHEREOF, the Parties have executed this Agreement.").style.font.bold = True

        sig_table = document.add_table(rows=5, cols=2)
        party1_name = parties[0].get("name", "Party A") if len(parties) > 0 else "First Party"
        party2_name = parties[1].get("name", "Party B") if len(parties) > 1 else "Second Party"

        sig_table.rows[0].cells[0].text = f"For: {party1_name}"
        sig_table.rows[0].cells[1].text = f"For: {party2_name}"
        sig_table.rows[1].cells[0].text = "Signature: ______________________"
        sig_table.rows[1].cells[1].text = "Signature: ______________________"
        sig_table.rows[2].cells[0].text = "Name: _________________________"
        sig_table.rows[2].cells[1].text = "Name: _________________________"
        sig_table.rows[3].cells[0].text = "Title: __________________________"
        sig_table.rows[3].cells[1].text = "Title: __________________________"
        sig_table.rows[4].cells[0].text = "Date: __________________________"
        sig_table.rows[4].cells[1].text = "Date: __________________________"

        # Footer
        footer = document.sections[0].footer
        footer_p = footer.paragraphs[0]
        footer_p.text = f"{branding.get('footer_text', 'Generated via LegalEase AI')} • DISCLAIMER: Informational draft only; not legal advice."
        footer_p.style.font.size = Pt(8)
        footer_p.style.font.color.rgb = RGBColor(148, 163, 184)

        buffer = io.BytesIO()
        document.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    @staticmethod
    def generate_txt(doc: Document) -> str:
        """Generates clean, readable plain text export."""
        branding: Dict[str, Any] = doc.branding_config or {}
        structured: Dict[str, Any] = doc.structured_content or {}
        sections: List[Dict[str, Any]] = structured.get("sections", [])
        terms: List[Dict[str, Any]] = structured.get("important_terms", [])
        parties: List[Dict[str, Any]] = structured.get("parties", [])

        lines = []
        org_name = branding.get("org_name", "LEGALEASE AI")
        lines.append("=" * 80)
        lines.append(f"{org_name.center(80)}")
        lines.append(f"{doc.title.upper().center(80)}")
        lines.append("=" * 80)
        lines.append(f"Document Type  : {doc.document_type.replace('_', ' ').title()}")
        lines.append(f"Effective Date : {structured.get('effective_date', 'As Executed')}")
        lines.append("-" * 80)

        if terms:
            lines.append("\nSUMMARY OF KEY CONTRACT TERMS:")
            lines.append("-" * 40)
            for t in terms:
                lines.append(f"  * {t.get('term', '')}: {t.get('value', '')}")
            lines.append("-" * 40)

        lines.append("\nAGREEMENT CLAUSES:\n")
        for sec in sorted(sections, key=lambda s: s.get("order", 0)):
            lines.append(sec.get("heading", "").upper())
            lines.append("~" * len(sec.get("heading", "")))
            lines.append(sec.get("content", ""))
            lines.append("")

        lines.append("\n" + "=" * 80)
        lines.append("IN WITNESS WHEREOF, the Parties have executed this Agreement:\n")

        p1 = parties[0].get("name", "Party A") if len(parties) > 0 else "First Party"
        p2 = parties[1].get("name", "Party B") if len(parties) > 1 else "Second Party"

        lines.append(f"FOR: {p1:<35} FOR: {p2}")
        lines.append(f"By:   _________________________        By:   _________________________")
        lines.append(f"Name: _________________________        Name: _________________________")
        lines.append(f"Date: _________________________        Date: _________________________")
        lines.append("\n" + "-" * 80)
        lines.append("LEGAL DISCLAIMER:")
        lines.append(structured.get("disclaimer", (
            "LegalEase provides AI-generated document drafts for informational and drafting purposes. "
            "These documents do not constitute legal advice and should be reviewed by a qualified legal professional before use."
        )))
        lines.append("=" * 80)

        return "\n".join(lines)


export_service = ExportService()
